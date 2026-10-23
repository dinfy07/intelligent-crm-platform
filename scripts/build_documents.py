from pathlib import Path
import json,re,shutil,os
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1]
M=json.loads((R/'docs/content_manifest.json').read_text())
OUT=Path(os.environ.get('CRM_DOCUMENT_OUTPUT',R/'deliverables'));OUT.mkdir(exist_ok=True)
def base():
 d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=Inches(.7);s.bottom_margin=Inches(.7);s.left_margin=Inches(.75);s.right_margin=Inches(.75)
 for name,size in [('Normal',11),('Title',25),('Subtitle',13),('Heading 1',19),('Heading 2',13),('Heading 3',11.5)]:
  st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
  if name.startswith('Heading'):st.font.bold=True
  for k in ['ascii','hAnsi','eastAsia','cs']:st.element.get_or_add_rPr().rFonts.set(qn('w:'+k),'Arial')
  for e in st.element.xpath('.//w:pBdr'):e.getparent().remove(e)
 n=d.styles['Normal'].paragraph_format;n.space_after=Pt(4);n.line_spacing=1.08
 for name in ['Heading 1','Heading 2','Heading 3']:
  f=d.styles[name].paragraph_format;f.space_before=Pt(10);f.space_after=Pt(5);f.keep_with_next=True
 d.styles['Title'].paragraph_format.space_after=Pt(15)
 p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 run=p.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
 d.core_properties.author='Project team';d.core_properties.title=M['title'];return d

def clean(t,strip=True):
 t=t.replace('<br>','\n').replace('**','').replace('`','')
 t=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',t)
 return t.strip() if strip else t
def addp(d,t,style=None):
 p=d.add_paragraph(style=style);p.paragraph_format.widow_control=True
 if t.startswith('**Priority:**'):p.paragraph_format.keep_with_next=True
 # Preserve bold runs where explicitly used in prose.
 for i,part in enumerate(t.split('**')):p.add_run(clean(part,False)).bold=(i%2==1)
 return p

def table(d,headers,rows,widths=None,size=10):
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths is None:
  width=(d.sections[-1].page_width-d.sections[-1].left_margin-d.sections[-1].right_margin)/914400
  if len(headers)==2:widths=[width*.28,width*.72]
  else:widths=[width/len(headers)]*len(headers)
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=clean(h)
 hp=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');hp.append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for i,val in enumerate(row[:len(headers)]):cells[i].text=clean(str(val))
 for ri,row in enumerate(t.rows):
  tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    v=OxmlElement('w:'+side);v.set(qn('w:val'),'single');v.set(qn('w:sz'),'5');v.set(qn('w:color'),'D9D9D9');b.append(v)
   pr.append(b);mar=OxmlElement('w:tcMar')
   for side,val in [('top','80'),('bottom','80'),('left','100'),('right','100')]:
    v=OxmlElement('w:'+side);v.set(qn('w:w'),val);v.set(qn('w:type'),'dxa');mar.append(v)
   pr.append(mar)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9EEF5');pr.append(sh)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
    for r in p.runs:r.font.size=Pt(size);r.bold=(ri==0)
 d.add_paragraph().paragraph_format.space_after=Pt(1)
 return t

def md(d,path,skip_title=True):
 lines=path.read_text().splitlines();i=0;first=True
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  if line.startswith('# '):
   if not skip_title:addp(d,clean(line[2:]),'Heading 1')
   i+=1;continue
  if line.startswith('## '):
   if line=='## Вопросы комиссии':d.add_page_break()
   addp(d,clean(line[3:]),'Heading 2');i+=1;continue
  if line.startswith('### '):addp(d,clean(line[4:]),'Heading 3');i+=1;continue
  if line.startswith('|'):
   block=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    block.append([v.strip() for v in lines[i].strip().strip('|').split('|')]);i+=1
   if len(block)>1 and all(set(c.replace(':','').replace(' ',''))<=set('-') for c in block[1]):block.pop(1)
   headers=block[0];rows=block[1:]
   if headers==['Aspect','Critical analysis']:
    for a,b in rows:addp(d,f'**{a}:** {b}')
   elif headers and headers[0]=='Decision' and len(headers)==5:
    for row in rows:
     addp(d,row[0],'Heading 2')
     for label,val in zip(headers[1:],row[1:]):addp(d,f'**{label}:** {val}')
   elif headers and headers[0]=='ID' and 'Acceptance criterion' in headers:
    for row in rows:
     item=dict(zip(headers,row));addp(d,f"{item['ID']} {item['Requirement']}",'Heading 3')
     addp(d,f"**Priority:** {item['Priority']}"+(f". **Owner:** {item['Owner']}" if 'Owner' in item else ''))
     addp(d,item['Acceptance criterion']);addp(d,'**Verification:** '+item['Verification'])
   else:
    widths=[1.45,2.05,2.05,2.05,2.05] if len(headers)==5 and headers[0]=='Criterion' else None
    table(d,headers,rows,widths=widths,size=9.5 if len(headers)>=5 else 10)
   continue
  if line.startswith('```'):
   code=[];i+=1
   while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
   p=addp(d,'\n'.join(code));p.paragraph_format.left_indent=Inches(.15)
   for r in p.runs:r.font.name='Courier New';r.font.size=Pt(9)
   i+=1;continue
  if re.match(r'^\d+\.\s',line) or line.startswith('- '):
   addp(d,line);i+=1;continue
  para=[line];i+=1
  while i<len(lines) and lines[i].strip() and not re.match(r'^(#|\||```|\d+\.\s|- )',lines[i].strip()):para.append(lines[i].strip());i+=1
  addp(d,' '.join(para))

def orient(d,landscape):
 s=d.add_section(WD_SECTION_START.NEW_PAGE)
 s.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
 s.page_width=Inches(11 if landscape else 8.5);s.page_height=Inches(8.5 if landscape else 11)
 s.left_margin=Inches(.65 if landscape else .75);s.right_margin=s.left_margin;s.top_margin=Inches(.65 if landscape else .7);s.bottom_margin=s.top_margin

d=base();p=d.add_paragraph('Intelligent CRM\nFirst pre defense package','Title');p.paragraph_format.space_before=Pt(36)
addp(d,M['title'],'Subtitle');addp(d,'Software Engineering · SE-2426\nAstana IT University\nCheckpoint: 24 October 2026\nPrepared: 19 October 2026; tracker/status updated: 23 October 2026')
table(d,['Team member','Responsibility'],M['team'],size=11)
addp(d,'Supervisor: Arailym Orazova')
addp(d,'This package proposes a text-first CRM and a reproducible evaluation. It includes a runnable interaction prototype and analysis/design materials. The source documents distinguish published evidence, proposed design choices and measurements that have not yet been made.')
addp(d,'**Readiness:** User discovery and supervisor feedback are not collected. YouTrack is configured with 14 assigned issues. Actual student contribution and peer-review evidence remain open. The package is a preparation candidate, not a claim of complete checkpoint compliance.')
addp(d,'Repository: https://github.com/dinfy07/intelligent-crm-platform\nPreparation branch: codex/CRM-13-pre-defense-package (AI-assisted review candidate; students commit their own meaningful revisions)')
d.add_page_break();addp(d,'Contents and readiness','Heading 1')
chapters=[('01_proposal.md','1 Project proposal'),('02_relevance_and_discovery.md','2 Relevance evidence and discovery'),('03_literature_review.md','3 Critical literature review'),('04_existing_systems.md','4 Existing system comparison'),('05_technology_comparison.md','5 Technology comparison'),('06_requirements.md','6 Initial requirements'),('07_architecture.md','7 Architecture and diagrams'),('08_ux_prototype.md','8 UX prototype'),('09_verification_evaluation.md','9 Verification and evaluation plan'),('10_management_risks.md','10 Project management and risks'),('11_supervisor_feedback.md','11 Supervisor feedback'),('12_sources.md','12 Source register')]
for _,h in chapters:addp(d,h)
addp(d,'Evidence checklist','Heading 2')
table(d,['Component','Current state'],[[r[1],r[2]] for r in M['checklist']],size=10)
prevland=False
for file,title in chapters:
 land=file.startswith('04_')
 if land!=prevland:orient(d,land)
 else:d.add_page_break()
 prevland=land;addp(d,title,'Heading 1')
 if file.startswith('12_'):
  addp(d,'All sources checked on 19 October 2026. Scientific sources are distinguished from living product and technical documentation. Product descriptions are not independent quality measurements.')
  sources=json.loads((R/'docs/sources.json').read_text())
  for source in sources:
   p=addp(d,f"**{source['id']} {source['authors']} ({source.get('year') or 'living documentation'}).** {source['title']}. {source.get('venue','')}\n{source['url']}\nUse: {source['used_in']}. "+source.get('limitation',source.get('claim','')))
   p.paragraph_format.space_after=Pt(8);p.paragraph_format.keep_together=True
   for run in p.runs:run.font.size=Pt(10.5)
 else:
  start=len(d.paragraphs);md(d,R/'docs'/file)
  if file.startswith(('02_','06_')):
   for paragraph in d.paragraphs[start:]:
    paragraph.paragraph_format.line_spacing=1.01
    paragraph.paragraph_format.space_after=Pt(2)
    if paragraph.style.name.startswith('Heading'):paragraph.paragraph_format.space_before=Pt(8)
 if file.startswith('07_'):
  for n,caption in [('context','System context'),('use-cases','Use case diagram'),('components','Draft component diagram')]:
   d.add_page_break();addp(d,caption,'Heading 2');d.add_paragraph().add_run().add_picture(str(R/f'diagrams/exports/{n}.png'),width=Inches(7));addp(d,f'Editable source: diagrams/src/{n}.puml. Editable SVG: diagrams/exports/{n}.svg. Proposed design, not deployed components.')
 if file.startswith('08_'):
  d.add_page_break();addp(d,'Runnable prototype view','Heading 2');d.add_paragraph().add_run().add_picture(str(R/'prototype/preview-desktop.png'),width=Inches(7));addp(d,'The screenshot shows fictional records and a deterministic extractive baseline. Use prototype/index.html for the full interactive walkthrough.')
# Remove title-border defaults globally.
for s in d.styles:
 for e in s.element.xpath('.//w:pBdr'):e.getparent().remove(e)
d.save(OUT/'CRM_Predefense_1_Report_EN.docx')
g=base();addp(g,'Подготовка к защите CRM проекта','Title');addp(g,'Пояснения на русском и список действий до 24 октября 2026 года','Subtitle')
md(g,R/'docs/RU_defense_notes.md')
g.add_page_break();addp(g,'Что ещё нужно сделать до сдачи','Heading 1')
for title,body in [
('Интервью с пользователями','Провести 3–5 содержательных бесед с представителями будущих пользователей или получить реальное согласование объёма с заказчиком. Вопросы и форма записи находятся в docs/evidence. Сохранять только разрешённые обезличенные сведения. Ни один такой разговор пока не проведён.'),
('YouTrack','Проект создан: crm-se2426.youtrack.cloud. Задачи CRM-1–CRM-14 назначены трём участникам. Даниалу и Ахмедияру нужно принять приглашения и проверить доступ. Связывайте реальные PR с задачами и меняйте статусы по фактическому результату.'),
('Реальный вклад каждого','Каждый участник должен прочитать свои материалы, внести осмысленные изменения со своего аккаунта и выполнить ревью работы другого участника. Исходная подготовка выполнена с помощью ИИ и не доказывает длительную историю работы студентов.'),
('Отзыв руководителя','Показать предложение, архитектуру, прототип и план оценки. Записать реальные замечания, решения по ним и ссылки на исправления. Форма журнала не является согласованием.'),
('Финальная версия','Проверить доступ руководителя и комиссии, закрыть обязательные замечания, выбрать проверенный коммит и создать неизменяемый тег pre-defense-1. Не переносить уже сданный тег на другой коммит.'),
('Предлагаемый порядок выступления','Beybit: проблема, пользователи, UX и демонстрация. Danial: требования, технологии, архитектура и интеграции. Akhmediyar: научный обзор, методы анализа, QA и оценка. Распределение можно уточнить на репетиции без изменения технических ролей.')]:
 addp(g,title,'Heading 2');addp(g,body)
addp(g,'Документы на английском','Heading 2');addp(g,'CRM_Predefense_1_Report_EN.docx — полный пакет. CRM_Predefense_1_EN.pptx — презентация. Исходные Markdown, диаграммы, тесты и прототип включены в общий пакет для проверки через PR по CRM-13. Ссылка на репозиторий: https://github.com/dinfy07/intelligent-crm-platform. На 23 октября репозиторий публичный; доступ на запись есть у Ахмедияра, приглашение Даниала ожидает принятия.')
for s in g.styles:
 for e in s.element.xpath('.//w:pBdr'):e.getparent().remove(e)
g.save(OUT/'CRM_Defense_Guide_RU.docx')
print('Created English report and Russian defence guide')
